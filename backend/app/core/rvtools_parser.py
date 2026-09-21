"""Streaming RVTools workbook parser. Pure: no DB, no FastAPI.

The workbook is opened with openpyxl in **read-only, values-only** mode, so
rows stream off the zip member and the sheet is never materialized. The
only state held across the parse is a small per-VM facts dict built from
the auxiliary sheets (a few hundred bytes per VM), which is what lets the
``vInfo`` pass attach disks/snapshots/NICs without a second look-up pass.

Sheets read:

  ========== =============================================================
  vInfo      one row per VM — identity, placement, sizing (required)
  vDisk      per-disk mode / sharing / RDM / controller / datastore
  vSnapshot  snapshots present at export time
  vNetwork   per-NIC portgroup / adapter / MAC / connected
  vCPU       CPU hot-add
  vMemory    memory hot-add
  ========== =============================================================

Everything except vInfo is optional; ``ParseStats.sheets_found`` records
what was there so the assessment can say "no vDisk sheet — disk-level
blockers not evaluated" instead of implying a clean bill of health.

Header matching is case- and punctuation-insensitive (``_norm``) and
alias-based, because customer exports vary by RVTools version and by
whatever a spreadsheet round-trip did to them. The alias lists carry both
the RVTools 4.x names ("VI SDK Server", "Network #1", "Provisioned MiB")
and the simplified names this repo's older fixtures and CSV template use
("vCenter", "vSphere Networks", "Provisioned MB").
"""

from __future__ import annotations

import csv
import re
from collections.abc import Callable, Iterator
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

FACTS_VERSION = 1

# VM.name column width. Names are identity — never truncate one to fit.
MAX_VM_NAME = 255

_NORM_RE = re.compile(r"[^a-z0-9]")


def _norm(header: Any) -> str:
    return _NORM_RE.sub("", str(header).lower())


# field → accepted normalized headers, first match wins.
VINFO_ALIASES: dict[str, tuple[str, ...]] = {
    "name": ("vm", "name", "vmname"),
    "source_hostname": ("dnsname", "hostname", "sourcehostname", "fqdn"),
    "ip_address": ("primaryipaddress", "primaryip", "ipaddress", "ip"),
    "guest_os_config": (
        "osaccordingtotheconfigurationfile",
        "guestosfullname",
        "guestos",
        "osfamily",
        "os",
    ),
    "guest_os_tools": ("osaccordingtothevmwaretools",),
    "role": ("role", "tag"),
    "annotation": ("annotation", "notes", "comment", "comments"),
    "vcenter": ("visdkserver", "vcenter", "vcenterserver", "vcserver", "vc"),
    "cluster": ("cluster", "vspherecluster", "computecluster"),
    "folder": ("folder", "folderpath", "vmfolder", "vspherefolder"),
    "datacenter": ("datacenter",),
    "host": ("host", "esxihost"),
    "power_state": ("powerstate",),
    "template": ("template",),
    "srm_placeholder": ("srmplaceholder",),
    "cpus": ("cpus", "numcpu", "vcpus"),
    "memory": ("memory", "memorymb", "memorymib"),
    "nics": ("nics",),
    "disks": ("disks",),
    "provisioned": ("provisionedmib", "provisionedmb"),
    "in_use": ("inusemib", "inusemb"),
    "firmware": ("firmware",),
    "hw_version": ("hwversion",),
    "cbt": ("cbt",),
    "ft_state": ("ftstate",),
    "secure_boot": ("efisecureboot",),
    "consolidation_needed": ("consolidationneeded",),
    "enable_uuid": ("enableuuid",),
    "connection_state": ("connectionstate",),
    "guest_state": ("gueststate",),
    "resource_pool": ("resourcepool",),
    "vapp": ("vapp",),
    "cluster_rules": ("clusterrulenames", "clusterrules"),
    "passthru_hotplug": ("fixedpassthruhotplug",),
    "latency_sensitivity": ("latencysensitivity",),
    "path": ("path",),
    "moref": ("vmid", "moref", "vmmoref"),
    "vm_uuid": ("vmuuid", "instanceuuid"),
    "smbios_uuid": ("smbiosuuid", "biosuuid"),
    "networks": ("vspherenetworks", "networks", "portgroups", "portgroup", "network"),
    "datastores": ("vspheredatastores", "datastores", "datastore"),
}

# Operator-defined vSphere custom attributes RVTools exports as their own
# columns. Tier 2 of environment detection reads "Environment".
CUSTOM_ATTRIBUTE_KEYS = ("Environment", "App", "Application", "Tier", "Owner")

_OS_FAMILY_RULES: tuple[tuple[tuple[str, ...], str], ...] = (
    (("red hat", "rhel"), "rhel"),
    (("ubuntu",), "ubuntu"),
    (("centos",), "centos"),
    (("rocky",), "rocky"),
    (("alma",), "almalinux"),
    (("suse", "sles"), "sles"),
    (("debian",), "debian"),
    (("windows",), "windows"),
    (("oracle",), "oracle"),
)

_DATASTORE_IN_PATH_RE = re.compile(r"^\s*\[([^\]]+)\]")
# "Owner: x | App: y | Env: prod" — the packed annotation convention some
# sites use instead of real custom attributes.
_ANNOTATION_PAIR_RE = re.compile(r"(?:^|[|;\n])\s*([A-Za-z][A-Za-z ]{1,24}?)\s*[:=]\s*([^|;\n]+)")
_ANNOTATION_KEY_MAP = {
    "env": "Environment",
    "environment": "Environment",
    "app": "App",
    "application": "App",
    "tier": "Tier",
    "owner": "Owner",
}


@dataclass
class RowIssue:
    sheet: str
    row_number: int
    vm_name: str | None
    reason: str
    severity: str = "rejected"  # "rejected" | "warning"


@dataclass
class ParsedVM:
    row_number: int
    vcenter_hostname: str | None
    # Column values ready to set on ``app.models.vm.VM``.
    fields: dict[str, Any]


@dataclass
class ParseStats:
    sheets_found: list[str] = field(default_factory=list)
    rows_total: int | None = None
    rows_read: int = 0
    current_sheet: str | None = None


@dataclass
class ScanResult:
    detected_vcenters: dict[str, int]
    vm_rows: int
    sheets_found: list[str]


class RVToolsParseError(Exception):
    """The file can't be treated as an inventory export at all."""


# --------------------------------------------------------------------------
# Value coercion
# --------------------------------------------------------------------------
def _s(v: Any, max_len: int | None = None) -> str | None:
    if v is None:
        return None
    out = str(v).strip()
    if not out:
        return None
    return out[:max_len] if max_len else out


def _i(v: Any) -> int | None:
    if v is None or v == "":
        return None
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return int(v)
    try:
        return int(float(str(v).replace(",", "").strip()))
    except ValueError:
        return None


def _b(v: Any) -> bool | None:
    if v is None or v == "":
        return None
    if isinstance(v, bool):
        return v
    text = str(v).strip().lower()
    if text in ("true", "yes", "1", "enabled"):
        return True
    if text in ("false", "no", "0", "disabled"):
        return False
    return None


def _split_list(v: Any) -> list[str]:
    if v is None:
        return []
    return [p.strip() for p in re.split(r"[;,]", str(v)) if p.strip()]


def normalize_hostname(raw: Any) -> str | None:
    text = _s(raw)
    return text.lower().rstrip(".") or None if text else None


def shorten_os_family(raw: str | None) -> str | None:
    if not raw:
        return None
    lower = raw.lower()
    for needles, family in _OS_FAMILY_RULES:
        if any(n in lower for n in needles):
            return family
    return raw[:32]


def _datastore_from_path(path: Any) -> str | None:
    m = _DATASTORE_IN_PATH_RE.match(str(path)) if path else None
    return m.group(1).strip() if m else None


# --------------------------------------------------------------------------
# Sheet streaming
# --------------------------------------------------------------------------
class _Source:
    """Uniform row stream over an .xlsx workbook or a single-sheet CSV."""

    def __init__(self, path: Path):
        self.path = path
        self.is_csv = path.suffix.lower() == ".csv"
        self._wb = None
        if not self.is_csv:
            try:
                from openpyxl import load_workbook

                self._wb = load_workbook(path, read_only=True, data_only=True)
            except Exception as e:  # zipfile.BadZipFile, InvalidFileException, …
                raise RVToolsParseError(
                    "Not a readable .xlsx workbook. Legacy .xls exports must be "
                    f"re-saved as .xlsx. ({type(e).__name__})"
                ) from e

    def close(self) -> None:
        if self._wb is not None:
            self._wb.close()

    def sheet_names(self) -> dict[str, str]:
        """normalized name → actual name."""
        if self.is_csv:
            return {"vinfo": "vInfo"}
        return {_norm(n): n for n in self._wb.sheetnames}

    def declared_rows(self, actual_name: str) -> int | None:
        if self.is_csv:
            return None
        n = self._wb[actual_name].max_row
        return n - 1 if n and n > 1 else None

    def rows(self, actual_name: str) -> Iterator[tuple[int, dict[str, Any]]]:
        """Yield (1-based sheet row number, {normalized header: value})."""
        if self.is_csv:
            with self.path.open(newline="", encoding="utf-8-sig", errors="replace") as fh:
                yield from self._dict_rows(csv.reader(fh))
            return
        ws = self._wb[actual_name]
        # Writers other than Excel often emit a wrong <dimension> tag, and
        # read-only mode trusts it. Resetting makes openpyxl scan to the
        # real end of the sheet instead of stopping early.
        ws.reset_dimensions()
        yield from self._dict_rows(ws.iter_rows(values_only=True))

    @staticmethod
    def _dict_rows(raw_rows) -> Iterator[tuple[int, dict[str, Any]]]:
        headers: list[str] | None = None
        for idx, row in enumerate(raw_rows, start=1):
            if headers is None:
                if not row or all(c is None or str(c).strip() == "" for c in row):
                    continue
                headers = [_norm(c) if c is not None else "" for c in row]
                continue
            if row is None or all(c is None or str(c).strip() == "" for c in row):
                continue
            yield idx, {h: row[i] for i, h in enumerate(headers) if h and i < len(row)}


def _get(row: dict[str, Any], aliases: tuple[str, ...]) -> Any:
    for a in aliases:
        v = row.get(a)
        if v is not None and str(v).strip() != "":
            return v
    return None


def _vm_key(vcenter: str | None, moref: str | None, name: str | None) -> tuple[str, str]:
    """Join key between vInfo and the auxiliary sheets. MoRef when the
    export has it; name otherwise (both scoped by vCenter)."""
    return (vcenter or "", f"id:{moref}" if moref else f"name:{name or ''}")


# --------------------------------------------------------------------------
# Auxiliary sheets → per-VM facts
# --------------------------------------------------------------------------
_AUX_VCENTER = VINFO_ALIASES["vcenter"]
_AUX_MOREF = VINFO_ALIASES["moref"]
_AUX_NAME = ("vm", "name", "vmname")


def _aux_key(row: dict[str, Any]) -> tuple[str, str]:
    return _vm_key(
        normalize_hostname(_get(row, _AUX_VCENTER)),
        _s(_get(row, _AUX_MOREF)),
        _s(_get(row, _AUX_NAME)),
    )


def _disk_fact(row: dict[str, Any]) -> dict[str, Any]:
    path = _s(row.get("path"), 512)
    return {
        "label": _s(row.get("disk") or row.get("label"), 128),
        "capacity_mib": _i(row.get("capacitymib") or row.get("capacitymb")),
        "mode": _s(row.get("diskmode"), 64),
        "sharing": _s(row.get("sharingmode"), 64),
        "raw": _b(row.get("raw")),
        "raw_compat_mode": _s(row.get("rawcompmode"), 64),
        "controller": _s(row.get("controller"), 128),
        "shared_bus": _s(row.get("sharedbus"), 64),
        "thin": _b(row.get("thin")),
        "path": path,
        "datastore": _datastore_from_path(path),
    }


def _snapshot_fact(row: dict[str, Any]) -> dict[str, Any]:
    return {
        "name": _s(row.get("name"), 255),
        "created": _s(row.get("datetime"), 64),
        "size_mib": _i(row.get("sizemibtotal") or row.get("sizemibvmsn")),
        "quiesced": _b(row.get("quiesced")),
    }


def _nic_fact(row: dict[str, Any]) -> dict[str, Any]:
    return {
        "label": _s(row.get("niclabel") or row.get("label"), 128),
        "adapter": _s(row.get("adapter"), 64),
        "network": _s(row.get("network"), 255),
        "switch": _s(row.get("switch"), 255),
        "connected": _b(row.get("connected")),
        "mac": _s(row.get("macaddress"), 32),
        "ipv4": _s(row.get("ipv4address"), 255),
    }


def _load_aux(
    src: _Source,
    sheets: dict[str, str],
    stats: ParseStats,
    tick: Callable[[], None],
) -> dict[tuple[str, str], dict[str, Any]]:
    facts: dict[tuple[str, str], dict[str, Any]] = {}

    def bucket(row: dict[str, Any]) -> dict[str, Any]:
        return facts.setdefault(_aux_key(row), {})

    list_sheets = (
        ("vdisk", "disks", _disk_fact),
        ("vsnapshot", "snapshots", _snapshot_fact),
        ("vnetwork", "nics", _nic_fact),
    )
    for norm_name, key, build in list_sheets:
        actual = sheets.get(norm_name)
        if not actual:
            continue
        stats.current_sheet = actual
        for _row_no, row in src.rows(actual):
            if _b(row.get("template")):
                continue
            bucket(row).setdefault(key, []).append(build(row))
            stats.rows_read += 1
            tick()

    for norm_name, key in (("vcpu", "cpu_hot_add"), ("vmemory", "mem_hot_add")):
        actual = sheets.get(norm_name)
        if not actual:
            continue
        stats.current_sheet = actual
        for _row_no, row in src.rows(actual):
            bucket(row)[key] = _b(row.get("hotadd"))
            stats.rows_read += 1
            tick()
    return facts


# --------------------------------------------------------------------------
# vInfo → ParsedVM
# --------------------------------------------------------------------------
def _custom_attributes(row: dict[str, Any], annotation: str | None) -> dict[str, str]:
    out: dict[str, str] = {}
    for key in CUSTOM_ATTRIBUTE_KEYS:
        lower = key.lower()
        v = _get(row, (lower, f"customattribute{lower}", f"ca{lower}", f"custom{lower}"))
        if v is not None:
            out[key] = str(v).strip()[:256]
    if annotation:
        for raw_key, raw_val in _ANNOTATION_PAIR_RE.findall(annotation):
            mapped = _ANNOTATION_KEY_MAP.get(raw_key.strip().lower())
            if mapped and mapped not in out:
                out[mapped] = raw_val.strip()[:256]
    if "Application" in out and "App" not in out:
        out["App"] = out["Application"]
    return out


def _build_vm(
    row_number: int,
    row: dict[str, Any],
    aux: dict[tuple[str, str], dict[str, Any]],
    sheet: str,
) -> ParsedVM | RowIssue | list:
    name = _s(_get(row, VINFO_ALIASES["name"]))
    hostname = _s(_get(row, VINFO_ALIASES["source_hostname"]), 255)
    if not name and hostname:
        name = hostname.split(".")[0]
    if not name:
        return RowIssue(sheet, row_number, None, "missing VM name")
    if len(name) > MAX_VM_NAME:
        return RowIssue(
            sheet,
            row_number,
            name[:80] + "…",
            f"VM name is {len(name)} characters; the limit is {MAX_VM_NAME}. "
            "Names are identity and are never truncated — shorten it in vSphere.",
        )
    if _b(_get(row, VINFO_ALIASES["template"])):
        return RowIssue(sheet, row_number, name, "is a template, not a VM — skipped")
    if _b(_get(row, VINFO_ALIASES["srm_placeholder"])):
        return RowIssue(sheet, row_number, name, "is an SRM placeholder — skipped")

    vcenter = normalize_hostname(_get(row, VINFO_ALIASES["vcenter"]))
    moref = _s(_get(row, VINFO_ALIASES["moref"]), 64)
    facts_aux = aux.get(_vm_key(vcenter, moref, name)) or aux.get(_vm_key(vcenter, None, name), {})

    disks = facts_aux.get("disks", [])
    nics = facts_aux.get("nics", [])

    networks: list[str] = _split_list(_get(row, VINFO_ALIASES["networks"]))
    for n in range(1, 11):  # RVTools "Network #1" … "Network #8" (+headroom)
        v = _s(row.get(f"network{n}"))
        if v:
            networks.append(v)
    networks.extend(n["network"] for n in nics if n.get("network"))

    datastores: list[str] = _split_list(_get(row, VINFO_ALIASES["datastores"]))
    ds_from_path = _datastore_from_path(_get(row, VINFO_ALIASES["path"]))
    if ds_from_path:
        datastores.append(ds_from_path)
    datastores.extend(d["datastore"] for d in disks if d.get("datastore"))

    networks = sorted(set(networks))
    datastores = sorted(set(datastores))

    annotation = _s(_get(row, VINFO_ALIASES["annotation"]))
    attrs = _custom_attributes(row, annotation)
    guest_os = _s(_get(row, VINFO_ALIASES["guest_os_config"]), 255)
    tools_os = _s(_get(row, VINFO_ALIASES["guest_os_tools"]), 255)

    hardware_facts = {
        "v": FACTS_VERSION,
        "firmware": _s(_get(row, VINFO_ALIASES["firmware"]), 16),
        "hw_version": _s(_get(row, VINFO_ALIASES["hw_version"]), 16),
        "cbt": _b(_get(row, VINFO_ALIASES["cbt"])),
        "ft_state": _s(_get(row, VINFO_ALIASES["ft_state"]), 64),
        "secure_boot": _b(_get(row, VINFO_ALIASES["secure_boot"])),
        "consolidation_needed": _b(_get(row, VINFO_ALIASES["consolidation_needed"])),
        "enable_uuid": _b(_get(row, VINFO_ALIASES["enable_uuid"])),
        "connection_state": _s(_get(row, VINFO_ALIASES["connection_state"]), 32),
        "guest_state": _s(_get(row, VINFO_ALIASES["guest_state"]), 32),
        "guest_os_tools": tools_os,
        "smbios_uuid": _s(_get(row, VINFO_ALIASES["smbios_uuid"]), 64),
        "resource_pool": _s(_get(row, VINFO_ALIASES["resource_pool"]), 255),
        "vapp": _s(_get(row, VINFO_ALIASES["vapp"]), 255),
        "cluster_rules": _s(_get(row, VINFO_ALIASES["cluster_rules"]), 512),
        "passthru_hotplug": _b(_get(row, VINFO_ALIASES["passthru_hotplug"])),
        "latency_sensitivity": _s(_get(row, VINFO_ALIASES["latency_sensitivity"]), 32),
        "in_use_mb": _i(_get(row, VINFO_ALIASES["in_use"])),
        "cpu_hot_add": facts_aux.get("cpu_hot_add"),
        "mem_hot_add": facts_aux.get("mem_hot_add"),
        "disks": disks,
        "snapshots": facts_aux.get("snapshots", []),
        "nics": nics,
    }

    fields: dict[str, Any] = {
        "name": name,
        "source_hostname": hostname or name,
        "ip_address": _s(_get(row, VINFO_ALIASES["ip_address"]), 45),
        "os_family": shorten_os_family(guest_os or tools_os),
        "guest_os_full": guest_os or tools_os,
        "role": _s(_get(row, VINFO_ALIASES["role"]), 64),
        "notes": annotation[:1024] if annotation else None,
        "owner": (attrs.get("Owner") or "")[:128] or None,
        "application_hint": (attrs.get("App") or "")[:128] or None,
        "vsphere_networks": networks,
        "vsphere_datastores": datastores,
        "vsphere_cluster": _s(_get(row, VINFO_ALIASES["cluster"]), 255),
        "vsphere_folder": _s(_get(row, VINFO_ALIASES["folder"]), 512),
        "vsphere_datacenter": _s(_get(row, VINFO_ALIASES["datacenter"]), 255),
        "esxi_host": _s(_get(row, VINFO_ALIASES["host"]), 255),
        "power_state": _s(_get(row, VINFO_ALIASES["power_state"]), 32),
        "custom_attributes": attrs,
        "moref": moref,
        "vm_uuid": _s(_get(row, VINFO_ALIASES["vm_uuid"]), 64),
        "num_cpus": _i(_get(row, VINFO_ALIASES["cpus"])),
        "memory_mb": _i(_get(row, VINFO_ALIASES["memory"])),
        "disk_count": _i(_get(row, VINFO_ALIASES["disks"])) or (len(disks) or None),
        "nic_count": _i(_get(row, VINFO_ALIASES["nics"])) or (len(nics) or None),
        "provisioned_mb": _i(_get(row, VINFO_ALIASES["provisioned"])),
        "hardware_facts": hardware_facts,
    }

    out: list = [ParsedVM(row_number, vcenter, fields)]
    if not datastores:
        out.append(
            RowIssue(
                sheet,
                row_number,
                name,
                "imported, but no datastore could be determined — this VM "
                "cannot be storage-mapped until the export includes one",
                severity="warning",
            )
        )
    if not networks:
        out.append(
            RowIssue(
                sheet,
                row_number,
                name,
                "imported, but no network/portgroup listed — this VM will "
                "not appear in network mapping",
                severity="warning",
            )
        )
    return out


# --------------------------------------------------------------------------
# Public API
# --------------------------------------------------------------------------
def _vinfo_sheet(sheets: dict[str, str], src: _Source) -> str:
    if "vinfo" in sheets:
        return sheets["vinfo"]
    if src.is_csv:
        return "vInfo"
    # Single-sheet workbooks (a CSV template saved as .xlsx) are accepted;
    # a multi-sheet workbook without vInfo is not an RVTools export.
    if len(sheets) == 1:
        return next(iter(sheets.values()))
    raise RVToolsParseError(
        "No 'vInfo' sheet found. Export with RVTools 'Export all to Excel', "
        f"or upload a single-sheet file. Sheets present: {sorted(sheets.values())}"
    )


def scan(path: Path) -> ScanResult:
    """Cheap first pass: which vCenter hostnames does the file contain, and
    how many VM rows each? Drives the routing step before any write."""
    src = _Source(path)
    try:
        sheets = src.sheet_names()
        vinfo = _vinfo_sheet(sheets, src)
        counts: dict[str, int] = {}
        total = 0
        saw_name_column = False
        for _row_no, row in src.rows(vinfo):
            saw_name_column = saw_name_column or any(
                a in row for a in VINFO_ALIASES["name"] + VINFO_ALIASES["source_hostname"]
            )
            if _b(_get(row, VINFO_ALIASES["template"])):
                continue
            host = normalize_hostname(_get(row, VINFO_ALIASES["vcenter"])) or ""
            counts[host] = counts.get(host, 0) + 1
            total += 1
        if total and not saw_name_column:
            raise RVToolsParseError(
                "No VM name column found (expected 'VM', 'Name' or 'DNS Name')."
            )
        if total == 0:
            raise RVToolsParseError("The VM sheet has a header but no data rows.")
        found = [a for n, a in sheets.items() if n in _KNOWN_SHEETS]
        return ScanResult(counts, total, found)
    finally:
        src.close()


_KNOWN_SHEETS = ("vinfo", "vdisk", "vsnapshot", "vnetwork", "vcpu", "vmemory")


def iter_vms(
    path: Path,
    stats: ParseStats,
    on_progress: Callable[[ParseStats], None] | None = None,
    progress_every: int = 250,
) -> Iterator[ParsedVM | RowIssue]:
    """Stream the workbook as ParsedVM / RowIssue items.

    ``stats`` is mutated in place so the caller can persist progress;
    ``on_progress`` fires every ``progress_every`` source rows.
    """
    src = _Source(path)
    try:
        sheets = src.sheet_names()
        vinfo = _vinfo_sheet(sheets, src)
        stats.sheets_found = [a for n, a in sheets.items() if n in _KNOWN_SHEETS] or [vinfo]

        declared = [src.declared_rows(sheets[n]) for n in _KNOWN_SHEETS if n in sheets]
        stats.rows_total = (
            sum(declared) if declared and all(d is not None for d in declared) else None
        )

        def tick() -> None:
            if on_progress and stats.rows_read % progress_every == 0:
                on_progress(stats)

        aux = _load_aux(src, sheets, stats, tick)

        stats.current_sheet = vinfo
        if on_progress:
            on_progress(stats)
        for row_no, row in src.rows(vinfo):
            stats.rows_read += 1
            built = _build_vm(row_no, row, aux, vinfo)
            if isinstance(built, list):
                yield from built
            else:
                yield built
            tick()
        if on_progress:
            on_progress(stats)
    finally:
        src.close()
