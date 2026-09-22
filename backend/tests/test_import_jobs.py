"""Server-side RVTools ingestion: parser, job runner, API.

TestClient runs BackgroundTasks before returning the response, so a POST
returns with the scan (and, when routing was supplied, the import) already
finished — the tests read the terminal job state straight off the poll.
"""

from __future__ import annotations

import csv
import io
import json
import time
from collections import Counter
from pathlib import Path

import pytest
from openpyxl import Workbook
from sqlalchemy import select

from app.core.rvtools_parser import (
    ParsedVM,
    ParseStats,
    RowIssue,
    RVToolsParseError,
    iter_vms,
    scan,
)
from app.models.vm import VM

FIXTURE_1000 = Path(__file__).resolve().parents[2] / "tests" / "fixtures" / "rvtools-1000.xlsx"
FIXTURE_LEGACY = (
    Path(__file__).resolve().parents[2] / "tests" / "fixtures" / "sample-rvtools-v2.xlsx"
)
VCENTERS = [f"vc-site{i:02d}.hospital.example" for i in (1, 2, 3)]
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@pytest.fixture(autouse=True)
def _spool(tmp_path, monkeypatch):
    monkeypatch.setenv("IMPORT_SPOOL_DIR", str(tmp_path / "spool"))


def _register(client, hostname: str) -> int:
    r = client.post(
        "/api/sources/vcenters", json={"name": hostname.split(".")[0], "hostname": hostname}
    )
    assert r.status_code == 201, r.text
    return r.json()["id"]


def _upload(client, path_or_bytes, *, name="export.xlsx", **form):
    data = path_or_bytes if isinstance(path_or_bytes, bytes) else Path(path_or_bytes).read_bytes()
    return client.post(
        "/api/imports/rvtools",
        files={"file": (name, data, XLSX)},
        data={k: (json.dumps(v) if isinstance(v, dict) else str(v)) for k, v in form.items()},
    )


def _workbook(vinfo_rows: list[list], header: list[str], extra: dict | None = None) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "vInfo"
    ws.append(header)
    for r in vinfo_rows:
        ws.append(r)
    for title, rows in (extra or {}).items():
        sheet = wb.create_sheet(title)
        for r in rows:
            sheet.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# --------------------------------------------------------------------------
# Parser
# --------------------------------------------------------------------------
def test_scan_counts_vms_per_vcenter():
    result = scan(FIXTURE_1000)
    assert set(result.detected_vcenters) == set(VCENTERS)
    assert "vDisk" in result.sheets_found and "vSnapshot" in result.sheets_found
    # templates are not VMs and are excluded from the routing counts
    assert result.vm_rows == sum(result.detected_vcenters.values())


def test_parser_joins_aux_sheets_and_reports_issues():
    stats = ParseStats()
    items = list(iter_vms(FIXTURE_1000, stats))
    vms = [i for i in items if isinstance(i, ParsedVM)]
    issues = [i for i in items if isinstance(i, RowIssue)]

    assert stats.rows_total == stats.rows_read  # progress can reach 100%
    sev = Counter(i.severity for i in issues)
    assert sev == {"rejected": 3, "warning": 2}  # (in-file duplicate is caught by the importer)
    reasons = " | ".join(i.reason for i in issues)
    assert "missing VM name" in reasons and "template" in reasons and "255" in reasons

    assert all(v.fields["moref"] for v in vms)
    assert any(v.fields["name"] == "radiología-архив-画像-001" for v in vms)
    with_disks = [v for v in vms if v.fields["hardware_facts"]["disks"]]
    assert len(with_disks) > 990
    flat = [d for v in vms for d in v.fields["hardware_facts"]["disks"]]
    assert any(d["raw"] for d in flat)
    assert any(d["mode"] == "independent_persistent" for d in flat)
    assert any(d["sharing"] == "sharingMultiWriter" for d in flat)
    assert any(v.fields["hardware_facts"]["snapshots"] for v in vms)
    assert any(v.fields["hardware_facts"]["cbt"] is False for v in vms)
    # datastore is derived from the VMX/VMDK path, networks from Network #n + vNetwork
    sample = vms[0].fields
    assert sample["vsphere_datastores"] and sample["vsphere_networks"] and sample["esxi_host"]


def test_parser_reads_legacy_simplified_headers():
    """The repo's older 57-VM fixture uses 'vCenter' / 'vSphere Networks'."""
    items = list(iter_vms(FIXTURE_LEGACY, ParseStats()))
    vms = [i for i in items if isinstance(i, ParsedVM)]
    assert len(vms) == 57
    assert all(v.vcenter_hostname and v.fields["vsphere_networks"] for v in vms)
    # packed "Env: prod" annotation feeds custom attributes
    assert any(v.fields["custom_attributes"].get("Environment") for v in vms)


def test_parser_rejects_non_rvtools_workbook(tmp_path):
    p = tmp_path / "x.xlsx"
    p.write_bytes(_workbook([], ["a"], extra={"Other": [["b"]]}))
    wb = Workbook()
    wb.active.title = "Summary"
    wb.create_sheet("Charts")
    wb.save(p)
    with pytest.raises(RVToolsParseError, match="No 'vInfo' sheet"):
        scan(p)
    bad = tmp_path / "not-a-zip.xlsx"
    bad.write_bytes(b"hello")
    with pytest.raises(RVToolsParseError, match="Not a readable"):
        scan(bad)


# --------------------------------------------------------------------------
# API — routing flow
# --------------------------------------------------------------------------
def test_upload_without_routing_waits_then_imports(client):
    ids = {h: _register(client, h) for h in VCENTERS}
    r = _upload(client, FIXTURE_1000)
    assert r.status_code == 202, r.text
    job = client.get(f"/api/imports/{r.json()['id']}").json()
    assert job["status"] == "awaiting_mapping"
    detected = {d["hostname"]: d for d in job["detected_vcenters"]}
    assert set(detected) == set(VCENTERS)
    # registered hostnames are pre-suggested so the UI can pre-fill routing
    assert all(detected[h]["suggested_vcenter_id"] == ids[h] for h in VCENTERS)
    assert client.get("/api/vms/stats").json()["total"] == 0  # nothing written yet

    r = client.post(f"/api/imports/{job['id']}/start", json={"vcenter_mapping": ids})
    assert r.status_code == 202, r.text
    done = client.get(f"/api/imports/{job['id']}").json()
    assert done["status"] == "completed", done
    assert done["progress_percent"] == 100
    assert done["created_count"] == 1000


def test_unrouted_hostnames_are_skipped_and_counted(client):
    first = _register(client, VCENTERS[0])
    r = _upload(client, FIXTURE_1000, vcenter_mapping={VCENTERS[0]: first})
    job = client.get(f"/api/imports/{r.json()['id']}").json()
    assert job["status"] == "completed"
    skipped = {s["hostname"]: s["vm_count"] for s in job["result"]["skipped_unrouted"]}
    assert set(skipped) == set(VCENTERS[1:])
    assert job["created_count"] == 335  # 332 generated + 2 warning rows + 1 unicode name
    assert skipped[VCENTERS[2]] == 332


def test_unknown_vcenter_id_is_refused_before_spooling(client):
    r = _upload(client, FIXTURE_1000, default_vcenter_id=999)
    assert r.status_code == 409
    assert "unknown vCenter" in r.json()["detail"]


def test_xls_and_empty_uploads_are_refused(client):
    assert _upload(client, b"x", name="old.xls").status_code == 415
    assert _upload(client, b"x", name="notes.txt").status_code == 415
    assert _upload(client, b"", name="empty.xlsx").status_code == 422


def test_corrupt_workbook_fails_the_job_with_a_message(client):
    vc = _register(client, VCENTERS[0])
    r = _upload(client, b"this is not a zip", default_vcenter_id=vc)
    job = client.get(f"/api/imports/{r.json()['id']}").json()
    assert job["status"] == "failed"
    assert "Not a readable .xlsx" in job["error_message"]


# --------------------------------------------------------------------------
# The 1,000-VM acceptance path
# --------------------------------------------------------------------------
def test_1000_vm_import_is_idempotent_and_reports_rejects(client, db_session):
    ids = {h: _register(client, h) for h in VCENTERS}

    started = time.monotonic()
    job = client.get(
        f"/api/imports/{_upload(client, FIXTURE_1000, vcenter_mapping=ids).json()['id']}"
    ).json()
    elapsed = time.monotonic() - started
    assert job["status"] == "completed", job
    assert job["created_count"] == 1000
    assert job["updated_count"] == 0
    assert job["rows_rejected"] == 4 and job["rows_warned"] == 2
    assert elapsed < 60, f"1,000-VM import took {elapsed:.1f}s"

    # Environment detection ran on the server-side path (the regression
    # this phase fixes: RVTools imports used to land with environment=NULL).
    env = job["result"]["environment_distribution"]
    assert env.get("production", 0) > 600
    assert env.get("development", 0) > 50 and env.get("staging", 0) > 50
    assert env.get("unset", 0) < 50

    # Same NAME on two vCenters is legal.
    clones = db_session.scalars(select(VM).where(VM.name == "ehr-db-001")).all()
    assert len({v.source_vcenter_id for v in clones}) >= 2

    vm = db_session.scalars(select(VM).where(VM.moref.is_not(None))).first()
    assert vm.esxi_host and vm.num_cpus and vm.hardware_facts["v"] == 1
    assert vm.vsphere_cluster and vm.vsphere_folder  # no longer dropped at the schema boundary

    # Reject report: JSON + CSV
    rejects = client.get(f"/api/imports/{job['id']}/rejects?severity=rejected").json()
    assert rejects["total"] == 4
    assert any("duplicate of row" in i["reason"] for i in rejects["items"])
    body = client.get(f"/api/imports/{job['id']}/rejects.csv")
    assert body.headers["content-type"].startswith("text/csv")
    rows = list(csv.DictReader(io.StringIO(body.text)))
    assert len(rows) == 6 and {"sheet", "row", "vm_name", "severity", "reason"} <= set(rows[0])

    # Re-import: zero duplicates, zero updates.
    again = client.get(
        f"/api/imports/{_upload(client, FIXTURE_1000, vcenter_mapping=ids).json()['id']}"
    ).json()
    assert again["status"] == "completed"
    assert (again["created_count"], again["updated_count"]) == (0, 0)
    assert again["unchanged_count"] == 1000
    assert again["marked_missing_count"] == 0
    assert client.get("/api/vms/stats").json()["total"] == 1000


# --------------------------------------------------------------------------
# Upsert semantics
# --------------------------------------------------------------------------
HEADER = ["VM", "VM ID", "VI SDK Server", "CPUs", "Path", "Network #1", "Folder"]


def _row(name, moref, cpus=2, folder="/site/prod/app"):
    return [name, moref, "vc-a.example", cpus, f"[ds1] {name}/{name}.vmx", "vlan-10", folder]


def test_upsert_updates_renames_by_moref_and_marks_missing(client, db_session):
    vc = _register(client, "vc-a.example")
    first = _workbook(
        [_row("web-01", "vm-1"), _row("web-02", "vm-2"), _row("web-03", "vm-3")], HEADER
    )
    assert (
        client.get(
            f"/api/imports/{_upload(client, first, default_vcenter_id=vc).json()['id']}"
        ).json()["created_count"]
        == 3
    )

    # web-01 resized; web-02 renamed in vSphere (same MoRef); web-03 gone.
    second = _workbook([_row("web-01", "vm-1", cpus=8), _row("web-02-renamed", "vm-2")], HEADER)
    job = client.get(
        f"/api/imports/{_upload(client, second, default_vcenter_id=vc).json()['id']}"
    ).json()
    assert (job["created_count"], job["updated_count"], job["marked_missing_count"]) == (0, 2, 1)

    by_moref = {v.moref: v for v in db_session.scalars(select(VM)).all()}
    assert by_moref["vm-1"].num_cpus == 8
    assert by_moref["vm-2"].name == "web-02-renamed"
    assert by_moref["vm-3"].missing_from_last_upload is True


def test_create_only_never_touches_existing_rows(client, db_session):
    vc = _register(client, "vc-a.example")
    _upload(client, _workbook([_row("web-01", "vm-1")], HEADER), default_vcenter_id=vc)
    changed = _workbook([_row("web-01", "vm-1", cpus=32), _row("web-09", "vm-9")], HEADER)
    job = client.get(
        f"/api/imports/{_upload(client, changed, default_vcenter_id=vc, mode='create_only').json()['id']}"
    ).json()
    assert (job["created_count"], job["updated_count"], job["marked_missing_count"]) == (1, 0, 0)
    assert db_session.scalars(select(VM).where(VM.moref == "vm-1")).one().num_cpus == 2


def test_operator_set_environment_survives_reimport(client, db_session):
    vc = _register(client, "vc-a.example")
    wb = _workbook([_row("web-01", "vm-1", folder="/site/dev/app")], HEADER)
    _upload(client, wb, default_vcenter_id=vc)
    vm = db_session.scalars(select(VM)).one()
    assert (vm.environment, vm.environment_source) == ("development", "auto_detected")

    r = client.patch(f"/api/vms/{vm.id}/environment", json={"environment": "production"})
    assert r.status_code == 200, r.text
    _upload(client, wb, default_vcenter_id=vc)
    db_session.expire_all()
    vm = db_session.scalars(select(VM)).one()
    assert (vm.environment, vm.environment_source) == ("production", "user_set")


def test_thinner_reexport_does_not_blank_existing_data(client, db_session):
    vc = _register(client, "vc-a.example")
    _upload(client, _workbook([_row("web-01", "vm-1")], HEADER), default_vcenter_id=vc)
    thin = _workbook([["web-01", "vm-1", "vc-a.example"]], HEADER[:3])
    job = client.get(
        f"/api/imports/{_upload(client, thin, default_vcenter_id=vc).json()['id']}"
    ).json()
    assert job["updated_count"] == 0
    vm = db_session.scalars(select(VM)).one()
    assert vm.num_cpus == 2 and vm.vsphere_datastores == ["ds1"]


def test_csv_upload_is_parsed_server_side(client, db_session):
    vc = _register(client, "vc-a.example")
    text = "VM Name,Primary IP Address,Datastore,Network\nübung-01,10.0.0.9,ds1,vlan-10\n"
    r = _upload(client, text.encode(), name="inventory.csv", default_vcenter_id=vc)
    job = client.get(f"/api/imports/{r.json()['id']}").json()
    assert job["status"] == "completed" and job["created_count"] == 1
    assert db_session.scalars(select(VM)).one().name == "übung-01"


def test_csv_report_neutralizes_formula_names(client):
    vc = _register(client, "vc-a.example")
    # second row duplicates the first's MoRef → lands in the reject report
    wb = _workbook([_row("@SUM(1+1)", "vm-1"), _row("@SUM(1+1)", "vm-1")], HEADER)
    job_id = _upload(client, wb, default_vcenter_id=vc).json()["id"]
    text = client.get(f"/api/imports/{job_id}/rejects.csv").text
    assert "'@SUM(1+1)" in text


def test_second_upload_refused_while_one_is_active(client, db_session):
    from app.models.import_job import ImportJob

    db_session.add(ImportJob(id="busy", filename="a.xlsx", status="importing", actor="t"))
    db_session.commit()
    vc = _register(client, "vc-a.example")
    r = _upload(client, _workbook([_row("a", "vm-1")], HEADER), default_vcenter_id=vc)
    assert r.status_code == 409 and "still running" in r.json()["detail"]


def test_orphaned_jobs_are_failed_at_startup(client, db_session, engine):
    from sqlalchemy.orm import sessionmaker

    from app.core.import_jobs import fail_orphan_imports
    from app.models.import_job import ImportJob

    db_session.add(ImportJob(id="stuck", filename="a.xlsx", status="importing", actor="t"))
    db_session.commit()
    assert fail_orphan_imports(sessionmaker(bind=engine)) == 1
    db_session.expire_all()
    job = db_session.get(ImportJob, "stuck")
    assert job.status == "failed" and "idempotent" in job.error_message


# --------------------------------------------------------------------------
# Server-side filters over imported data
# --------------------------------------------------------------------------
def test_inventory_filters_on_imported_columns(client):
    vc = _register(client, "vc-a.example")
    header = [*HEADER, "Cluster", "Powerstate", "Primary IP Address"]
    rows = [
        [*_row("web-01", "vm-1"), "prod-cl01", "poweredOn", "10.1.1.10"],
        [*_row("web-02", "vm-2"), "prod-cl01", "poweredOff", "10.1.1.11"],
        [
            "db-01",
            "vm-3",
            "vc-a.example",
            4,
            "[tier_1] db-01/db-01.vmx",
            "vlan-100",
            "/s/prod",
            "dev-cl02",
            "poweredOn",
            "10.9.9.9",
        ],
    ]
    _upload(client, _workbook(rows, header), default_vcenter_id=vc)

    def names(qs):
        return sorted(v["name"] for v in client.get(f"/api/vms?{qs}").json()["items"])

    assert names("vsphere_cluster=prod-cl01") == ["web-01", "web-02"]
    assert names("power_state=poweredOff") == ["web-02"]
    # "vlan-10" must not match "vlan-100"; "_" in a datastore name is literal
    assert names("network=vlan-10") == ["web-01", "web-02"]
    assert names("datastore=tier_1") == ["db-01"]
    assert names("datastore=tierX1") == []
    assert names("search=10.9.9") == ["db-01"]  # IP search, as the UI placeholder promises

    facets = client.get("/api/vms/facets").json()
    assert facets["vsphere_cluster"] == {"prod-cl01": 2, "dev-cl02": 1}
    assert facets["power_state"] == {"poweredOn": 2, "poweredOff": 1}
    item = client.get("/api/vms?search=db-01").json()["items"][0]
    assert item["moref"] == "vm-3" and item["num_cpus"] == 4
